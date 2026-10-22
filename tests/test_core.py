# geodoc tests — plain python, zero frameworks; the referee is the seeded gold.
# Usage: python tests/test_core.py

import json
import os
import shutil
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "app"))

import llm  # noqa: E402
import store  # noqa: E402
import stages  # noqa: E402
from pdfpages import survey, has_text  # noqa: E402

import tempfile  # noqa: E402
import domain  # noqa: E402
# isolate the domain graph BEFORE any stage_outputs runs (its ingest hook
# writes through domain.DB)
domain.DB = os.path.join(tempfile.mkdtemp(prefix="geodoc-domain-"), "graph.db")

# A fresh clone has no 18 GB model and no llama.cpp build. The harness needs
# no LLM, but the KV state contract and the seal read the model manifest and
# the GGUF's size — so when the configured files are absent, the harness
# points llm at a stub manifest in a temp dir (same roles, ports and ctx as
# the shipped one; the server sits in a b0000/ folder so the seal still
# records a build id). With the real files present nothing is stubbed.
_mf = llm.manifest()
if not (os.path.isfile(_mf["server"])
        and all(os.path.isfile(c["gguf"]) for c in _mf["models"].values())):
    _stub = tempfile.mkdtemp(prefix="geodoc-stub-")
    os.makedirs(os.path.join(_stub, "b0000"))
    with open(os.path.join(_stub, "b0000", "llama-server.exe"), "wb") as _f:
        _f.write(b"\0")
    with open(os.path.join(_stub, "stub.gguf"), "wb") as _f:
        _f.write(b"\0" * (1 << 20))
    _shipped = json.load(open(llm.manifest_path(), encoding="utf-8"))
    _shipped["server"] = os.path.join(_stub, "b0000", "llama-server.exe")
    for _c in _shipped["models"].values():
        _c["gguf"] = os.path.join(_stub, "stub.gguf")
    llm.MANIFEST_LOCAL = os.path.join(_stub, "models.json")
    json.dump(_shipped, open(llm.MANIFEST_LOCAL, "w", encoding="utf-8"))
    print("note: model files absent; checks run against a stub manifest")

PDF = os.path.join(ROOT, "testdata", "operat_sample.pdf")
GOLD = json.load(open(os.path.join(ROOT, "testdata", "gold.json"), encoding="utf-8"))
FAILS = []
CHECKS = []


def check(name, cond, detail=""):
    print(("ok  " if cond else "FAIL") + f" {name}" + (f"  [{detail}]" if detail and not cond else ""))
    CHECKS.append(name)
    if not cond:
        FAILS.append(name)


# ---- pdfpages ----
sv = survey(PDF)
check("pdf: 5 pages", len(sv) == 5)
check("pdf: raster page flagged", [r["page"] for r in sv if not r["has_text"]] ==
      GOLD["strony_bez_tekstu"])
check("pdf: wykaz page has text", has_text(PDF, GOLD["strona_wykazu"]))

# ---- context store ----
cid = store.create(role="test", model="none", note="hello")
store.append(cid, "user", "pytanie 1")
store.append(cid, "assistant", "odpowiedz 1")
store.append(cid, "user", "pytanie 2")
fid = store.fork(cid, upto=3, note="fork-po-1-wymianie")
check("store: fork truncates", len(store.transcript(fid)) == 3 + 1)  # +1 fork note? note prepends
tr = store.transcript(fid)
check("store: fork keeps order", any(r["content"] == "odpowiedz 1" for r in tr))
check("store: parent recorded", store.meta(fid)["parent"].startswith(cid))
store.archive(cid)
store.archive(fid)
check("store: archive hides", all(m["id"] not in (cid, fid) for m in store.listing()))

# ---- prompt layering (golden/user overlay + slots) ----
pname = "klasyfikator"
golden_txt = llm.prompt(pname)
check("prompts: golden source by default", llm.prompt_info(pname)["source"] == "golden")
user_p = os.path.join(ROOT, "prompts", "user", pname + ".md")
os.makedirs(os.path.dirname(user_p), exist_ok=True)
with open(user_p, "w", encoding="utf-8") as f:
    f.write(golden_txt + "\n# NADPISANE\n")
check("prompts: user overlay wins", llm.prompt(pname).endswith("# NADPISANE\n")
      and llm.prompt_info(pname)["source"] == "user")
os.remove(user_p)
check("prompts: reset returns to golden", llm.prompt(pname) == golden_txt)
check("prompts: slot filled",
      llm.fill_slots("A\n{{PROFIL}}\nB", {"PROFIL": "hint"}) == "A\nhint\nB")
check("prompts: empty slot stripped", llm.fill_slots("A\n{{PROFIL}}\nB") == "A\nB")
check("prompts: all agents have golden manuals",
      llm.prompt_list() == ["detektor_rozdzialow", "ekstraktor_naglowka",
                            "ekstraktor_protokolu", "ekstraktor_wykazu",
                            "klasyfikator", "kompaktor", "opisywacz_zdjec",
                            "refleksja", "rewident", "syntezator_rozdzialow",
                            "werdyktor"])
check("prompts: detektor schema is GBNF-safe and loads",
      llm.schema("detektor_rozdzialow") is not None)
check("prompts: protokol scaffold has a schema (dormant until protocols)",
      llm.schema("ekstraktor_protokolu") is not None)

# ---- main chat window assembly (AutoCompact view) ----
import chat  # noqa: E402
rows_fake = [
    dict(role="note", content="start"),
    dict(role="user", content="u1"), dict(role="assistant", content="a1"),
    dict(role="user", content="u2"), dict(role="assistant", content="a2"),
    dict(role="note", kind="compact", content="## Decyzje\n- d1", keep_from=3),
    dict(role="user", content="u3"),
]
msgs = chat.build_messages(rows_fake)
check("chat: compact note folds into system + keeps tail",
      msgs[0]["role"] == "system" and "## Decyzje" in msgs[0]["content"]
      and [m["content"] for m in msgs[1:]] == ["u2", "a2", "u3"])
msgs2 = chat.build_messages(rows_fake[:5])
check("chat: no compact -> full window", [m["content"] for m in msgs2[1:]] ==
      ["u1", "a1", "u2", "a2"] and "## Decyzje" not in msgs2[0]["content"])
check("chat: kompaktor has schema", llm.schema("kompaktor") is not None)
ok1, allowed = llm.budget_check(9000, 3000, 16384)
ok2, _ = llm.budget_check(10000, 3000, 16384)
check("budget: 0.75 law as assertion", ok1 is True and allowed == 12288
      and ok2 is False)

# ---- KV state contract (slice 3) ----
c1 = llm.state_contract("worker")
check("state: contract stable and short", c1 == llm.state_contract("worker")
      and len(c1) == 12)
# rowne kontrakty miedzy rolami sa LEGALNE od wariantow jednomodelowych
# (Lekki/80B: obie role = ten sam gguf+ctx); inwariantem jest CZULOSC:
# kontrakt musi sie przesunac, gdy zmienia sie cokolwiek nosnego dla KV
_man_real = llm.manifest
_pert = _man_real()
_pert["models"]["worker"] = dict(_pert["models"]["worker"], ctx=8192)
llm.manifest = lambda: _pert
check("state: contract shifts when ctx changes",
      llm.state_contract("worker") != c1)
llm.manifest = _man_real
sd = llm.state_dir("worker")
check("state: dir is contract-versioned",
      sd.endswith(c1) and os.path.isdir(sd) and os.sep + "worker" + os.sep in sd)
cid_s = store.create(role="test-state", model="worker", note="state test")
st = store.attach_state(cid_s, "worker", c1, n_tokens=863)
check("state: attach recorded in meta", store.meta(cid_s)["state"]["contract"] == c1
      and st["n"] == 863)
store.archive(cid_s)

# ---- full pipeline on the synthetic operat ----
job_name = "test-run"
job_dir = os.path.join(ROOT, "jobs", job_name)
if os.path.isdir(job_dir):
    shutil.rmtree(job_dir)
res = stages.run_all(PDF, name=job_name)
hdr = json.load(open(os.path.join(job_dir, "extracted", "header.json"), encoding="utf-8"))
wyk = json.load(open(os.path.join(job_dir, "extracted", "wykaz.json"), encoding="utf-8"))
chk = res["checks"]

check("header: cel_pracy", hdr["cel_pracy"] == GOLD["cel_pracy"])
check("header: id_zgloszenia", hdr["id_zgloszenia"] == GOLD["id_zgloszenia"])
check("header: dzialka TERYT", hdr["dzialka"] == GOLD["dzialka"])
check("header: uklad", hdr["uklad"] == GOLD["uklad"])
check("header: nr_uprawnien", hdr["nr_uprawnien"] == GOLD["nr_uprawnien"])
check("header: daty", (hdr["data_pomiaru"], hdr["data_sporzadzenia"]) ==
      (GOLD["data_pomiaru"], GOLD["data_sporzadzenia"]))
check("header: rodzaj_aktu", hdr["rodzaj_aktu"] == GOLD["rodzaj_aktu"])
check("header: strona_wykazu", hdr["strona_wykazu"] == GOLD["strona_wykazu"])

check("wykaz: count", wyk["liczba_punktow"] == GOLD["liczba_punktow"])
gold_pts = {p["id"]: p for p in GOLD["punkty"]}
exact = sum(1 for p in wyk["punkty"]
            if p["id"] in gold_pts
            and abs(p["x"] - gold_pts[p["id"]]["x"]) < 0.005
            and abs(p["y"] - gold_pts[p["id"]]["y"]) < 0.005
            and p["status"] == gold_pts[p["id"]]["status"])
check("wykaz: every point exact to the cm + status", exact == len(gold_pts),
      f"{exact}/{len(gold_pts)}")
check("wykaz: attributes carried", all(
    (p["zrd"], p["bpp"], p["stb"], p["rzg"]) ==
    (GOLD["atrybuty"]["ZRD"], GOLD["atrybuty"]["BPP"],
     GOLD["atrybuty"]["STB"], GOLD["atrybuty"]["RZG"]) for p in wyk["punkty"]))

check("checks: PASS", chk["PASS"] is True, json.dumps(chk))
check("checks: shoelace vs gold area", abs(chk["pole_shoelace_m2"] - GOLD["pole_m2"]) < 0.6)

# outputs
import csv
p1, p2 = res["outputs"][:2]
rows = list(csv.reader(open(p1, encoding="utf-8-sig"), delimiter=";"))
check("out: przeglad has data row", len(rows) == 2 and rows[1][3] == GOLD["dzialka"])
rows2 = list(csv.reader(open(p2, encoding="utf-8-sig"), delimiter=";"))
check("out: upload rows = points+1", len(rows2) == GOLD["liczba_punktow"] + 1)
check("out: upload carries page refs", all(r[-1].isdigit() for r in rows2[1:]))
check("out: statuses preserved", {r[7] for r in rows2[1:]} ==
      {p["status"] for p in GOLD["punkty"]})

# ---- exports (slice 6): xlsx + per-PODGiK template ----
outs = [os.path.basename(x) for x in res["outputs"]]
check("out: xlsx produced", "przeglad.xlsx" in outs)
from openpyxl import load_workbook
wb = load_workbook(os.path.join(job_dir, "out", "przeglad.xlsx"))
check("out: xlsx sheets + rows",
      wb.sheetnames == ["Przegląd", "Wykaz punktów"]
      and wb["Wykaz punktów"].max_row == GOLD["liczba_punktow"] + 1)
_ws2 = wb["Wykaz punktów"]
_src_col = _ws2.max_column
check("out: xlsx ZRODLO column deep-links into the PDF",
      _ws2.cell(row=1, column=_src_col).value == "ŹRÓDŁO"
      and "#page=3" in (_ws2.cell(row=2, column=_src_col).hyperlink.target or ""))

# ---- rentgen operatu: value -> source geometry ----
from pdfpages import find_value_boxes
_gx = f"{GOLD['punkty'][0]['x']:.2f}"
_boxes = find_value_boxes(PDF, 3, _gx)
check("rentgen: value located on the wykaz page with sane geometry",
      len(_boxes) >= 1 and all(0 <= v <= 1 for b in _boxes for v in b)
      and _boxes[0][1] > 0.1)
check("rentgen: absent value -> no boxes",
      find_value_boxes(PDF, 3, "0000000.00") == [])
from pdfpages import render_page as _rp
import ocr as _ocr
_scan_png = os.path.join(ROOT, "jobs", job_name, "pages", "rentgen-p3.png")
_rp(os.path.join(ROOT, "testdata", "operat_scan.pdf"), 3, _scan_png)
_o3 = _ocr.ocr_page(_scan_png)
_hit = next((l for l in _o3["lines"] if "5571240.00" in (l.get("text") or "")
             and l.get("box")), None)
check("rentgen: OCR box lands where the coordinate actually sits",
      _hit is not None and all(0 <= v <= 1 for v in _hit["box"])
      and 0.15 < _hit["box"][0] < 0.45 and 0.05 < _hit["box"][1] < 0.30)
os.remove(_scan_png)
tpl_dir = os.path.join(ROOT, "profiles")
os.makedirs(tpl_dir, exist_ok=True)
with open(os.path.join(tpl_dir, "podgik-testowy.json"), "w", encoding="utf-8") as f:
    json.dump(dict(kolumny=["ID_PKT", "Y", "X", "STB"], separator=","), f)
outs2 = stages.stage_outputs(job_name, template="testowy")
rows_t = list(csv.reader(open(outs2[1], encoding="utf-8-sig"), delimiter=","))
check("out: PODGiK template reorders columns + separator",
      rows_t[0] == ["ID_PKT", "Y", "X", "STB"]
      and rows_t[1][1].startswith("6") and rows_t[1][2].startswith("5"))
os.remove(os.path.join(tpl_dir, "podgik-testowy.json"))
stages.stage_outputs(job_name)  # restore default outputs for later checks

# transcript image reference for the raster page
md = open(os.path.join(job_dir, "transcript", "page-005.md"), encoding="utf-8").read()
check("transcript: raster page references its image", "../pages/page-005.png" in md)

# ---- document map (canonical document.md + docmap.json) ----
dm = json.load(open(os.path.join(job_dir, "docmap.json"), encoding="utf-8"))
check("mapa: folio sequence read", [p["folio"] for p in dm["pages"]] ==
      [1, 2, 3, 4, None])
check("mapa: raster page typed skan + real image coverage",
      dm["pages"][4]["typ"] == "skan" and dm["pages"][4]["pokrycie_obrazem"] > 0.9)
check("mapa: wykaz page typed tabela", dm["pages"][2]["typ"] == "tabela")
check("mapa: no false gaps", dm["braki"]["folio_luki"] == [] and
      dm["braki"]["fragment_od"] is None)
check("mapa: toc covers all pages", [t["page"] for t in dm["toc"]] == [1, 2, 3, 4, 5])
doc_md = open(os.path.join(job_dir, "document.md"), encoding="utf-8").read()
check("mapa: document.md anchors + pdf backlink",
      "{#p003}" in doc_md and 'source/operat.pdf#page=3' in doc_md)
check("mapa: document.md front matter + toc",
      doc_md.startswith("---\ndoc_id: test-run\n") and "## Spis treści" in doc_md)
check("mapa: raster section carries image ref",
      "![strona 5](pages/page-005.png)" in doc_md)
g, frag, _ = stages.folio_gaps([(1, 1), (2, 2), (3, 4), (4, None), (5, 5)])
check("mapa: gap detector finds true removal",
      len(g) == 1 and g[0]["brakuje"] == 1 and g[0]["zakres"] == [3, 3]
      and frag is None)
g2, frag2, _ = stages.folio_gaps([(1, 121), (2, 122), (3, 125)])
check("mapa: fragment start + inner gaps",
      frag2 == 121 and g2[0]["brakuje"] == 2 and g2[0]["zakres"] == [123, 124])
g3, _, _ = stages.folio_gaps([(3, 3), (4, None), (5, None), (6, 6), (7, 7)])
check("mapa: unnumbered-but-present pages are NOT gaps (unnumbered-middle case)", g3 == [])
_, _, n4 = stages.folio_gaps([(1, 5), (2, 3)])
check("mapa: backwards folio -> note, not gap", len(n4) == 1 and "niespójna" in n4[0])
noise = stages.folio_consensus([(1, 121), (2, 121), (3, None), (4, 59), (5, 9)])
check("mapa: consensus rejects duplicated + inconsistent folios (repeated case-number noise)",
      noise == [None] * 5)
real = stages.folio_consensus([(1, 121), (2, 122), (3, None), (4, 124), (5, 999)])
check("mapa: consensus keeps arithmetic sequence, drops outlier",
      real == [121, 122, None, 124, None])

# ---- OCR branch (slice 4a): RapidOCR digit fidelity on our own print ----
import ocr
check("ocr: RapidOCR available", ocr.available())
_ocr_png = os.path.join(ROOT, "jobs", job_name, "pages", "ocr-probe-p3.png")
from pdfpages import render_page
render_page(PDF, 3, _ocr_png)
_o = ocr.ocr_page(_ocr_png)
check("ocr: reads the wykaz print confidently",
      _o["mean_conf"] is not None and _o["mean_conf"] > 0.9
      and len(_o["lines"]) > 8)
_coords = [f"{p['x']:.2f}" for p in GOLD["punkty"]] + \
          [f"{p['y']:.2f}" for p in GOLD["punkty"]]
_found = sum(1 for c in _coords if c in _o["text"])
check("ocr: >=16/18 coordinates digit-exact (200dpi clean print)",
      _found >= 16, f"{_found}/18")
os.remove(_ocr_png)
agr = ocr.coord_agreement("a 5571240.00 b 6553238.24", "5571240.00 x 9999999.99")
check("ocr: G4 agreement classifies both/one-engine coords",
      agr["obie"] == ["5571240.00"] and agr["tylko_pierwszy"] == ["6553238.24"]
      and agr["tylko_drugi"] == ["9999999.99"])

# ---- gates: G5 vocabulary, ranges, G3 referee, G2 grounding ----
check("g5: legal attributes pass",
      stages.check_slownik([dict(id="a", zrd="1", bpp="6", stb="4", rzg="7")])
      == (True, []))
s_old, viol = stages.check_slownik(
    [dict(id="a", zrd="1", bpp="9", stb="1", rzg="1")], "2019-05-01")
check("g5: violation on pre-2021 operat is a hard fail",
      s_old is False and viol == ["a: bpp=9"])
s_new, _ = stages.check_slownik(
    [dict(id="a", zrd="1", bpp="9", stb="1", rzg="1")], "2022-01-01")
check("g5: violation on 2021+ operat is a review-None", s_new is None)
check("g5: ranges tightened", stages.X_RANGE == (5_300_000, 6_130_000)
      and stages.ZONE_Y["6"] == (6_290_000, 6_710_000))

rows = stages.parse_rows(
    "152.4-01     5571240.00    6553238.24    1   1   3   1  bez_zmian\n"
    "152.4-02     5571274.79    6553221.47    1   1   3   1  bez_zmian\n")
check("g3: parse_rows shared parser", len(rows) == 2
      and rows[0]["x"] == 5571240.00 and rows[1]["id"] == "152.4-02")
ref = stages.referee_wykaz(
    {3: rows},
    [dict(id="152.4-01", x=5571240.00, y=6553238.24, zrd="1", bpp="1",
          stb="3", rzg="1", page=3),
     dict(id="152.4-99", x=1.0, y=2.0, zrd=None, bpp=None, stb=None,
          rzg=None, page=3)])
check("g3: referee flags omission, tolerates llm-only rows",
      ref["status"] is False and ref["zgodne"] == 1 and ref["tylko_llm"] == 1
      and ref["pominiete_przez_llm"] == [dict(page=3, id="152.4-02")])
ref2 = stages.referee_wykaz(
    {3: rows[:1]},
    [dict(id="152.4-01", x=5571240.10, y=6553238.24, zrd="1", bpp="1",
          stb="3", rzg="1", page=3)])
check("g3: referee catches a 10 cm transcription drift",
      ref2["status"] is False and len(ref2["niezgodne"]) == 1)
check("g3: nothing to referee on scans -> None",
      stages.referee_wykaz({}, [])["status"] is None)

g = stages.grounding_check(
    PDF,
    dict(id_zgloszenia="GK.6640.1234.2026", data_pomiaru="2026-06-11",
         uklad="PL-2000/6", zmyslone="TO-NIE-ISTNIEJE-9Z"),
    dict(id_zgloszenia=1, data_pomiaru=1, uklad=1, zmyslone=2))
check("g2: real values ground on claimed page",
      g["pola"]["id_zgloszenia"]["claim_ok"] is True
      and g["pola"]["uklad"]["found_page"] == 1)
check("g2: invented value = hallucination fail",
      g["pola"]["zmyslone"]["found_page"] is None
      and g["wartosci_w_dokumencie"] is False)
g2 = stages.grounding_check(PDF, dict(id_zgloszenia="GK.6640.1234.2026"),
                            dict(id_zgloszenia=3))
check("g2: wrong source page = mis-attribution flag, value still found",
      g2["pola"]["id_zgloszenia"]["found_page"] == 1
      and g2["pola"]["id_zgloszenia"]["claim_ok"] is False
      and g2["wartosci_w_dokumencie"] is True and g2["zrodla_trafne"] is False)

# ---- wave-5 gates: diacritic fold, dup-coord flag, ctx ceiling law ----
# model 'corrects' diacritics absent from source: scan page 5 OCR says
# "geo krakow" (no diacritics), model transcribes "GEO Kraków" — tier-2 fold
# grounds it and routes to arbitration instead of a hallucination FAIL
g3 = stages.grounding_check(
    PDF, dict(wykonawca="GEO Kraków"), dict(wykonawca=5),
    ocr_texts={5: "wykonawca geo krakow sp z o o"})
check("w5: diacritic-only mismatch folds, grounds and flags for arbitration",
      g3["pola"]["wykonawca"]["found_page"] == 5
      and g3["pola"]["wykonawca"].get("fold") is True
      and g3["wartosci_w_dokumencie"] is True
      and g3["diakrytyka_fold"] == ["wykonawca"])
check("w5: exact matches never carry the fold flag",
      "fold" not in g["pola"]["id_zgloszenia"]
      and g["diakrytyka_fold"] is None)
check("w5: fold helper strips PL diacritics incl. l-stroke",
      stages._fold_diacritics("Żółć łąka ćma") == "Zolc laka cma")
check("w5: ctx law passes under ceiling",
      llm.ctx_law(dict(ctx=16384, ctx_ceiling=65536)) == 16384)
check("w5: dup-coord flag catches distinct ids on one cm-exact position",
      stages.dup_coords([
          dict(id="A", x=1.0, y=2.0), dict(id="B", x=1.0, y=2.0),
          dict(id="C", x=3.0, y=4.0)]) == [["A", "B"]]
      and stages.dup_coords([dict(id="A", x=1.0, y=2.0)]) == [])
try:
    llm.ctx_law(dict(ctx=131072, ctx_ceiling=65536))
    check("w5: ctx over ceiling raises (DeltaNet g_cum law)", False)
except AssertionError:
    check("w5: ctx over ceiling raises (DeltaNet g_cum law)", True)

# ---- routing resilience: classifier under-nomination cannot lose pages ----
import agents as _ag
_row = "152.4-%02d   5571240.10   6553238.24   1   2   3   1   istniejący"
_texts = {1: "SPRAWOZDANIE TECHNICZNE bez tabel",
          2: "\n".join(_row % i for i in range(1, 10)),
          3: "\n".join(_row % i for i in range(10, 19)),
          4: "\n".join(_row % i for i in range(19, 28))}
check("routing: code detection unions pages the classifier missed",
      _ag.routed_wykaz_pages([2], _texts) == [2, 3, 4])
check("routing: no false wykaz on prose pages",
      _ag.routed_wykaz_pages([], {1: _texts[1]}) == [])
check("routing: single stray row is not a wykaz page",
      _ag.routed_wykaz_pages([], {5: _row % 1}) == [])

# ---- rozdzialy (code layer) vs the stress-fixture gold ----
import rozdzialy as _rz
_gs = json.load(open(os.path.join(ROOT, "testdata", "gold_stress.json"),
                     encoding="utf-8"))
_ch = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata",
                                       "operat_chaptered.pdf"))
_gold_ch = [(c["od"], c["do"]) for c in _gs["operat_chaptered"]["chapters"]]
check("rozdzialy: chaptered — 5 chapters at exact gold boundaries",
      [(c["od"], c["do"]) for c in _ch["rozdzialy"]] == _gold_ch,
      str(_ch["rozdzialy"]))
check("rozdzialy: chaptered — TOC found and entries parsed",
      _ch["spis_tresci"] and _ch["spis_tresci"]["strona_spisu"] == 2
      and len(_ch["spis_tresci"]["wpisy"]) == 5)
check("rozdzialy: chaptered — coverage gates green",
      _ch["pokrycie_ciagle"] and _ch["do_konca"]
      and _ch["braki_numeracji"] == [])
_ci = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata",
                                       "operat_incomplete.pdf"))
_gold_ci = [(c["od"], c["do"]) for c in _gs["operat_incomplete"]["chapters"]]
check("rozdzialy: incomplete — 4 real chapters, none invented",
      [(c["od"], c["do"]) for c in _ci["rozdzialy"]] == _gold_ci,
      str(_ci["rozdzialy"]))
check("rozdzialy: incomplete — numbering gap FLAGGED not fabricated",
      _ci["braki_numeracji"] and "brak rozdzialu 4" in _ci["braki_numeracji"][0]
      and _ci["spis_tresci"] is None)
# adversary family: lying TOC, contamination, degenerate, massive
_tk = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata",
                                       "operat_toc_klamie.pdf"))
check("adwersarz: lying TOC — rogue killed by legality, drift recovered",
      [(c["od"], c["do"]) for c in _tk["rozdzialy"]]
      == [(c["od"], c["do"]) for c in _gs["operat_toc_klamie"]["chapters"]]
      and _tk["spis_zgodny"] is False and _tk["spis_przesuniecie"] == 2)
_ob2 = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata", "operat_obcy.pdf"))
check("adwersarz: foreign page absorbed without breaking tiling",
      [(c["od"], c["do"]) for c in _ob2["rozdzialy"]]
      == [(c["od"], c["do"]) for c in _gs["operat_obcy"]["chapters"]]
      and _ob2["podzial_kompletny"])
_kr = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata", "operat_krotki.pdf"))
check("adwersarz: heading-free doc yields ZERO chapters, all nieoznaczone",
      _kr["rozdzialy"] == [] and _kr["nieoznaczone"] == [[1, 3]]
      and _kr["podzial_kompletny"])
_ms = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata",
                                       "operat_masywny.pdf"))
check("adwersarz: masywny 122 pages — 6 chapters exact, spis zgodny",
      [(c["od"], c["do"]) for c in _ms["rozdzialy"]]
      == [(c["od"], c["do"]) for c in _gs["operat_masywny"]["chapters"]]
      and _ms["spis_zgodny"] is True)

_scan_page = 3        # chapter-1 heading page of the pageocr twin
from pdfpages import render_page as _rp3
import ocr as _ocr3
_png3 = _rp3(os.path.join(ROOT, "testdata", "operat_pageocr.pdf"),
             _scan_page, os.path.join(ROOT, "jobs", "_rz-ocr-probe.png"))
_ob = {_scan_page: _ocr3.ocr_page(_png3)["lines"]}
_cands_scan = _rz.heading_candidates(
    os.path.join(ROOT, "testdata", "operat_pageocr.pdf"), ocr_boxes=_ob)
check("rozdzialy: scan page yields the chapter heading via OCR boxes",
      any(c["page"] == _scan_page and "SPRAWOZDANIE" in c["text"].upper()
          for c in _cands_scan),
      str([c["text"] for c in _cands_scan if c["page"] == _scan_page][:4]))

# ---- walker chain (fake runner) + assembly union + G11/G12 + folders ----
import walker as _wk
_chap_pdf = os.path.join(ROOT, "testdata", "operat_chaptered.pdf")
_rzjob = stages.run_all(_chap_pdf, name="rz-e2e")["job"]

_script = iter([
    dict(nastepny_start=3, nastepny_tytul="1. SPRAWOZDANIE TECHNICZNE",
         pewnosc="wysoka"),
    dict(nastepny_start=99, nastepny_tytul="ZMYSLONY", pewnosc="wysoka"),
    dict(nastepny_start=12, nastepny_tytul="4. PROTOKOLY GRANICZNE",
         pewnosc="wysoka"),
    dict(koniec_dokumentu=True),
])
_claims = _wk.walk(_rzjob, _chap_pdf,
                   runner=lambda payload: next(_script, dict(
                       koniec_dokumentu=True)))
check("walker: containment rejects a claim outside code candidates",
      any(c["pewnosc"] == "odrzucona-przez-kod" for c in _claims))
check("walker: claims file written, chain monotone and terminates",
      os.path.isfile(os.path.join(ROOT, "jobs", _rzjob, "extracted",
                                  "chapters_claims.json"))
      and all(a["okno"][0] < b["okno"][0]
              for a, b in zip(_claims, _claims[1:])))
check("walker: confirmed_claims filters to high-confidence code-contained",
      {c["nastepny_start"] for c in _wk.confirmed_claims(_claims)} == {3, 12})

# union: strip anchors from candidates -> walker claims must carry the spine
_cands_plain = [dict(page=3, text="Pierwszy", num=None, size=15, level=1),
                dict(page=9, text="Drugi", num=None, size=15, level=1)]
_asm = _rz.chapters_assemble(
    _cands_plain, None, 16,
    claims=[dict(nastepny_start=9, nastepny_tytul="Drugi", pewnosc="wysoka")])
check("union: walker claim anchors the spine when code has no anchors",
      len(_asm["rozdzialy"]) == 1 and _asm["rozdzialy"][0]["od"] == 9
      and _asm["rozdzialy"][0]["zrodlo"] == "walker"
      and _asm["strony_wstepne"] == 8 and _asm["nieoznaczone"] == [])
_asm2 = _rz.chapters_assemble(_cands_plain, None, 16, claims=None)
check("union: G11 tiling holds with and without claims",
      _asm["podzial_kompletny"] and _asm2["podzial_kompletny"])

# e2e phase-A on the chaptered job: assemble + G12 + materialize + docmap
_doc = _rz.rozdzialy_job(_rzjob, _chap_pdf)
check("rozdzialy-job: 5 chapters, G11+G12 green, provenance recorded",
      len(_doc["rozdzialy"]) == 5 and _doc["podzial_kompletny"]
      and _doc["tytuly_ugruntowane"]
      and len(_doc["prowieniencja_stron"]) >= 14)
_rdir = os.path.join(ROOT, "jobs", _rzjob, "rozdzialy")
check("rozdzialy-job: folders + master.md with linked TOC",
      os.path.isfile(os.path.join(_rdir, "master.md"))
      and sum(1 for d in os.listdir(_rdir)
              if os.path.isdir(os.path.join(_rdir, d))) == 5
      and "rozdzial.md)" in open(os.path.join(_rdir, "master.md"),
                                 encoding="utf-8").read())
check("rozdzialy-job: docmap carries the chapters key",
      "chapters" in json.load(open(os.path.join(ROOT, "jobs", _rzjob,
                                                "docmap.json"),
                                   encoding="utf-8")))
_g12 = _rz.guard_titles(
    dict(rozdzialy=[dict(nr=1, tytul="NIE MA TAKIEGO TYTULU", od=3, do=5,
                         num=None, zrodlo="kod", spis=None)],
         nieoznaczone=[]),
    {3: "zupelnie inna tresc strony"})
check("g12: ungrounded title demoted to nieoznaczone, not kept green",
      _g12["rozdzialy"] == [] and _g12["nieoznaczone"] == [[3, 5]]
      and _g12["tytuly_ugruntowane"] is False)

# ---- uklady: GK math, detection signatures, conversion provenance ----
import math as _math
import uklady as _uk
_phi, _lam = _math.radians(52.2297), _math.radians(21.0122)
_x92, _y92 = _uk.gk_forward(_phi, _lam, _uk.PL1992)
_p2, _l2 = _uk.gk_inverse(_x92, _y92, _uk.PL1992)
check("uklady: GK roundtrip machine-clean (<0.1 mm)",
      6.4e6 * max(abs(_p2 - _phi), abs(_l2 - _lam)) < 1e-4)
_x20, _y20, _z = _uk.convert_1992_to_2000(_x92, _y92)
check("uklady: Warszawa 1992->2000 lands in zone 7 with sane magnitudes",
      _z == 7 and 5_700_000 < _x20 < 5_900_000
      and 7_400_000 < _y20 < 7_600_000)
_xb, _yb = _uk.convert_2000_zone(_x20, _y20, 7, 6)
_xc, _yc = _uk.convert_2000_zone(_xb, _yb, 6, 7)
check("uklady: zone-to-zone conversion closes within 1 cm",
      abs(_xc - _x20) <= 0.01 and abs(_yc - _y20) <= 0.01)
check("uklady: detection - declared beats absent, magnitudes cross-check",
      _uk.wykryj_uklad("PL-2000/6", [dict(x=5571240.1, y=6553238.2)])
      == dict(system="PL-2000", strefa=6, sygnatura="PL-2000",
              zgodnosc_deklaracji_z_liczbami=True)
      and _uk.wykryj_uklad(None, [dict(x=487654.3, y=637890.1)])["system"]
      == "PL-1992"
      and _uk.wykryj_uklad("uklad lokalny m. X",
                           [dict(x=5432.1, y=8765.4)])["system"] == "lokalny")
check("uklady: 1992 conversion adds derived columns, original untouched",
      (lambda pts, meta: pts[0]["x"] == 487654.3
       and "x_2000" in pts[0] and meta["przeliczenie"] is not None
       and meta["dokladnosc_m"] == 0.001)(
          *_uk.przelicz_punkty([dict(x=487654.3, y=637890.1)],
                               dict(system="PL-1992"))))
check("uklady: 1965 and lokalny ABSTAIN with reasons, never silent-convert",
      _uk.przelicz_punkty([dict(x=4_800_000.0, y=3_600_000.0)],
                          dict(system="PL-1965"))[1]["uwaga"] is not None
      and _uk.przelicz_punkty([dict(x=5432.1, y=8765.4)],
                              dict(system="lokalny"))[1]["uwaga"] is not None)

# ---- fala 8: pewnosc tiers + template rename/transpose + 1992 pipeline ----
_pw = stages.pewnosc_punktow(
    [dict(id="A", page=3), dict(id="B", page=3), dict(id="C", page=4)],
    dict(referee_zgodny=True, wartosci_w_dokumencie=True,
         flagi=dict(wsp_jeden_silnik=[dict(id="B", wsp="5571240.10")],
                    powtorzone_wspolrzedne=[["C", "A"]])))
check("pewnosc: tiers from real signals (green/amber), reasons carried",
      _pw["A"]["tier"] == "SPRAWDZ" and _pw["B"]["tier"] == "SPRAWDZ"
      and _pw["C"]["tier"] == "SPRAWDZ"
      and "jednym silnikiem" in _pw["B"]["powody"][0])
_pw2 = stages.pewnosc_punktow([dict(id="A", page=3)],
                              dict(referee_zgodny=False, flagi={}))
check("pewnosc: referee mismatch = KRYTYCZNE",
      _pw2["A"]["tier"] == "KRYTYCZNE")
_pw3 = stages.pewnosc_punktow([dict(id="A", page=3)],
                              dict(referee_zgodny=True, flagi={}))
check("pewnosc: clean point = PEWNE", _pw3["A"]["tier"] == "PEWNE")
_j92 = stages.run_all(os.path.join(ROOT, "testdata", "operat_1992.pdf"),
                      name="ars-1992")
_c92 = _j92["checks"]
check("fala8: 1992 detected, PL-2000 gates neutralized not failed",
      _c92["uklad_wykryty"] == "PL-1992"
      and _c92["strefa_vs_wspolrzedne"] is None
      and _c92["x_w_zakresie"] is None
      and "uklad_do_recenzji" in (_c92.get("flagi") or {}))
import csv as _csv
_csv_p = [os.path.join(ROOT, "jobs", "ars-1992", "out", f)
          for f in os.listdir(os.path.join(ROOT, "jobs", "ars-1992", "out"))
          if f.startswith("wgranie")][0]
_rows92 = list(_csv.reader(open(_csv_p, encoding="utf-8-sig"),
                           delimiter=";"))
check("fala8: 1992 upload CSV carries derived PL-2000 columns + originals",
      "X_2000" in _rows92[0] and "X" in _rows92[0]
      and abs(float(_rows92[1][_rows92[0].index("X_2000")])
              - 5_788_456) < 2_000)
_tpl = dict(stages.DEFAULT_TEMPLATE, kolumny=["ID_PKT", "X", "Y"],
            nazwy_kolumn={"ID_PKT": "NUMER"}, transpozycja=True)
stages.stage_outputs("ars-1992", template=_tpl)
_rows_t = list(_csv.reader(open(_csv_p, encoding="utf-8-sig"),
                           delimiter=";"))
check("fala8: operator template renames columns and transposes",
      _rows_t[0][0] == "NUMER" and len(_rows_t) <= 6
      and len(_rows_t[0]) >= 7)
stages.stage_outputs("ars-1992")      # restore default artifact

# ---- arsenal: scanner pathologies + abstention at scale ----
_pp = _rz.page_pathologies({
    1: "# Strona 1\ntresc pierwsza",
    2: "# Strona 2\n**BRAK WARSTWY TEKSTOWEJ** — x\n![strona 2](../p.png)\n"
       "Transkrypcja OCR (RapidOCR, x):\n```text\n\n```",
    3: "# Strona 3\ntresc druga",
    4: "# Strona 4\ntresc druga"})
check("arsenal: blank core detected as pusta, double-feed as dublet",
      _pp["puste"] == [2] and _pp["dublety"] == [[4, 3]])
_bn = _rz.stage_rozdzialy(os.path.join(ROOT, "testdata",
                                       "operat_bez_naglowkow.pdf"))
check("arsenal: 40 heading-less pages -> zero chapters, honest nieoznaczone",
      _bn["rozdzialy"] == [] and _bn["nieoznaczone"] == [[1, 40]]
      and _bn["podzial_kompletny"] and _bn["kandydaci"] == 0)

# ---- slice 3: opisywacz — media discovery, grounding, md patch ----
import opisy as _op
_med = _op.find_media(_chap_pdf)
check("opisy: chaptered media found with captions (2 maps + photo page)",
      len(_med) >= 3 and any("Mapa 1" in (m["caption"] or "") for m in _med)
      and all(m["page"] in (6, 7, 8) for m in _med))
_g = _op.ground_elements(
    dict(elementy=["strzalka polnocy", "smok wawelski"]),
    "mapa ze strzalka polnocy i skala")
check("opisy: ungrounded element dropped, grounded kept",
      _g["elementy"] == ["strzalka polnocy"]
      and _g["elementy_odrzucone"] == ["smok wawelski"]
      and _g["elementy_ugruntowane"] is False)
_fake_opis_calls = []
def _fake_opis(payload):
    _fake_opis_calls.append(payload)
    return dict(typ="mapa", opis="Szkic granic dzialki.",
                elementy=["szkic przebiegu granic"])
_om = _op.describe_media(_rzjob, _chap_pdf, runner=_fake_opis)
check("opisy: fresh call per image + media.json written",
      len(_fake_opis_calls) == len(_om) and len(_om) >= 3
      and os.path.isfile(os.path.join(ROOT, "jobs", _rzjob, "extracted",
                                      "media.json")))
_ch2md = open(os.path.join(ROOT, "jobs", _rzjob, "rozdzialy",
                           "02-2-szkice-i-mapy", "rozdzial.md"),
              encoding="utf-8").read()
check("opisy: rozdzial.md patched with grounded description block",
      "OPIS-MEDIA START" in _ch2md and "Szkic granic dzialki." in _ch2md)

# ---- slice 2: summaries with G13 grounding + excel (fake runner) ----
import podsumowania as _ps
def _fake_synth(payload):
    m = __import__("re").search(r"strony (\d+)-(\d+)", payload)
    a = int(m.group(1))
    return dict(typ_rozdzialu="inny", streszczenie="Testowe streszczenie.",
                fakty=[
                    dict(fakt="sredni blad 0.02 m", strona=a),
                    dict(fakt="kwota 9999.99 zl", strona=a),
                    dict(fakt="fakt spoza zakresu", strona=99)],
                braki=[])
_sums = _ps.summarize_chapters(_rzjob, runner=_fake_synth)
check("g13: in-range fact with verbatim number survives",
      all(len(s["fakty"]) == (1 if s["nr"] == 1 else 0) or s["nr"] != 1
          for s in _sums)
      and any(f["fakt"].startswith("sredni blad")
              for f in _sums[0]["fakty"]))
check("g13: wrong-number and out-of-range facts dropped + flagged",
      _sums[0]["fakty_ugruntowane"] is False
      and len(_sums[0].get("fakty_odrzucone") or []) == 2)
check("g13: trailing sentence punctuation never fails a true number",
      _ps._numbers("wspolrzedne X: 5571240.00, Y: 6553218.85.")
      == ["5571240.00", "6553218.85"]
      and _ps.g13_ground(
          dict(fakty=[dict(fakt="X: 5571240.00, Y: 6553218.85", strona=3)]),
          dict(od=3, do=5),
          {3: "wiersz 5571240.00   6553218.85 w transkrypcie"}
      )["fakty_ugruntowane"] is True)
_ph = _ps.phase_podsumowania(_rzjob, runner=_fake_synth)
check("slice2: excel with two sheets + master digest appended",
      os.path.isfile(_ph["excel"])
      and "Streszczenie dokumentu" in open(
          os.path.join(ROOT, "jobs", _rzjob, "rozdzialy", "master.md"),
          encoding="utf-8").read())

# ---- E9: residency decision (dual on big-RAM dev, single on client) ----
_mcfg = llm.manifest()["models"]
check("e9: 190 GB dev box decides dual",
      llm.residency_mode(_mcfg, 190 * 2**30) == "dual")
check("e9: 8 GB box decides single",
      llm.residency_mode(_mcfg, 8 * 2**30) == "single")
_wg = _mcfg["worker"]["gguf"]
check("e9: shared gguf counted once (mmap) -> dual on 64 GB",
      llm.residency_mode({"a": dict(gguf=_wg), "b": dict(gguf=_wg)},
                         64 * 2**30) == "dual")
check("e9: explicit override wins",
      llm.residency_mode(_mcfg, 190 * 2**30, override="single") == "single")

# ---- G10: digit fidelity floor on the seeded DEGRADED scan ----
_deg = os.path.join(ROOT, "testdata", "wykaz_degraded.png")
if os.path.isfile(_deg):
    _od = _ocr.ocr_page(_deg)
    _found_deg = sum(1 for c in _coords if c in _od["text"])
    check("g10: degraded-scan digit fidelity >= 16/18 (regression floor)",
          _found_deg >= 16, f"{_found_deg}/18")
else:
    check("g10: degraded fixture present (tools/degrade_scan.py)", False)

# ---- pieczec replay: deterministic reproducibility seal ----
seal1 = json.load(open(os.path.join(job_dir, "checks", "pieczec.json"),
                       encoding="utf-8"))
seal2 = stages.stage_pieczec(job_name)
check("pieczec: deterministic (two builds byte-identical)",
      seal1 == json.loads(json.dumps(seal2)))
check("pieczec: seals artifacts + engine build",
      "extracted/wykaz.json" in seal1["artefakty"]
      and "out/przeglad.csv" in seal1["artefakty"]
      and seal1["silnik"]["build"].startswith("b")
      and seal1["seed"] == 7420)
# the judge's verdict and the chapter map are sealed too (a hand-edited
# verdict or chapter must not keep a valid seal); the judge's prompt hash
# comes from verdict/, where its _meta lives
_vdir = os.path.join(job_dir, "verdict")
os.makedirs(_vdir, exist_ok=True)
with open(os.path.join(_vdir, "werdykt.json"), "w", encoding="utf-8") as f:
    json.dump(dict(_agent="werdyktor", werdykt="OK",
                   _meta=dict(prompt_hash="abc123", model="worker")), f)
_seal_v = stages.stage_pieczec(job_name)
_seal_rz = stages.stage_pieczec(_rzjob)
check("pieczec: covers verdict/ and the judge's prompt hash",
      "verdict/werdykt.json" in _seal_v["artefakty"]
      and _seal_v["agenci"].get("werdyktor", {}).get("prompt_hash") == "abc123")
check("pieczec: covers the chapter map recursively",
      "rozdzialy/master.md" in _seal_rz["artefakty"]
      and any(k.startswith("rozdzialy/") and k.count("/") >= 2
              for k in _seal_rz["artefakty"]))
shutil.rmtree(_vdir)
stages.stage_pieczec(job_name)
_wyk_p = os.path.join(job_dir, "extracted", "wykaz.json")
_orig = open(_wyk_p, encoding="utf-8").read()
with open(_wyk_p, "a", encoding="utf-8") as f:
    f.write(" ")
check("pieczec: hand-edit of an artifact changes its hash",
      stages.stage_pieczec(job_name)["artefakty"]["extracted/wykaz.json"]
      != seal1["artefakty"]["extracted/wykaz.json"])
with open(_wyk_p, "w", encoding="utf-8") as f:
    f.write(_orig)
stages.stage_pieczec(job_name)

# ---- supervised crew: state machine with a fake runner ----
import crew as _crew
_cjob = _crew.start(PDF, "crew-test")["job"]
_calls = []
def _fake_runner(job, st, dopisek):
    _calls.append((_crew.STAGES[st["cursor"]], dopisek))
    return dict(ok=True, _meta=dict(prompt_hash="deadbeef0000"))
r1 = _crew.step(_cjob, "run", runner=_fake_runner)
check("crew: run reports stage + awaits operator",
      r1["awaiting"] and r1["stage"] == "klasyfikator"
      and r1["prompt_hash"] == "deadbeef0000")
r2 = _crew.step(_cjob, "retry", dopisek="daty tylko ze sprawozdania",
                runner=_fake_runner)
_ov = open(os.path.join(ROOT, "jobs", _cjob, "tasks", "overrides.jsonl"),
           encoding="utf-8").read()
check("crew: retry records the override and re-runs the SAME stage",
      r2["stage"] == "klasyfikator" and "sprawozdania" in _ov
      and _calls[-1] == ("klasyfikator", "daty tylko ze sprawozdania"))
r3 = _crew.step(_cjob, "accept")
check("crew: accept advances the cursor", r3["next_stage"] == "naglowek")
for _ in range(3):
    _crew.step(_cjob, "accept")
check("crew: walk to done", _crew.step(_cjob, "accept")["done"] is True)
check("crew: dopisek block labels the operator authority",
      "DOPISEK OPERATORA" in _crew._dopisek_block("x")
      and _crew._dopisek_block(None) == "")
shutil.rmtree(os.path.join(ROOT, "jobs", "crew-test"), ignore_errors=True)

# ---- nocna zmiana: batch queue with a fake crew runner ----
import time
import nightshift as _ns
def _fake_crew(pdf):
    time.sleep(0.05)
    return dict(job="fake-job-" + os.path.basename(pdf)[:6], PASS=True,
                werdykt="OK", flagi=0)
_bid = _ns.submit([PDF, os.path.join(ROOT, "testdata", "nie-ma.pdf"), PDF],
                  runner=_fake_crew)
for _ in range(80):
    _st = _ns.status(_bid)
    if _st["done"]:
        break
    time.sleep(0.1)
check("nocna: batch completes with per-item statuses",
      _st["done"] and [i["status"] for i in _st["items"]] ==
      ["done", "blad", "done"])
check("nocna: missing file marked without killing the batch",
      _st["items"][1]["blad"] == "plik nie istnieje")
_rap = open(os.path.join(ROOT, "jobs", f"kolejka-{_bid}",
                         "raport_poranny.md"), encoding="utf-8").read()
check("nocna: raport poranny counts PASS and links documents",
      "Przetworzono 2/3" in _rap and "PASS: 2/2" in _rap
      and "document.md" in _rap and "BŁĄD" in _rap)
check("nocna: status exposes the raport link", _st["raport"] is not None
      and _st["raport"].endswith("raport_poranny.md"))
# audit find: a second batch submitted WHILE the worker lives must not be
# lost — the worker drains a queue of bids
_bid2 = _ns.submit([PDF], runner=_fake_crew)
_bid3 = _ns.submit([PDF], runner=_fake_crew)
for _ in range(100):
    if _ns.status(_bid2)["done"] and _ns.status(_bid3)["done"]:
        break
    time.sleep(0.1)
check("nocna: batch submitted during a live worker is NOT lost",
      _ns.status(_bid2)["done"] and _ns.status(_bid3)["done"])
for _b in (_bid2, _bid3):
    shutil.rmtree(os.path.join(ROOT, "jobs", f"kolejka-{_b}"),
                  ignore_errors=True)
shutil.rmtree(os.path.join(ROOT, "jobs", f"kolejka-{_bid}"),
              ignore_errors=True)

# ---- domain graph (slice 9): bi-temporal rodowod punktu ----
_attrs = dict(zrd="1", bpp="2", stb="1", rzg="1", status="istniejący", page=3)
_hdrA = dict(data_pomiaru="2020-05-10", wykonawca="GEO-A", dzialka="101/1")
_hdrB = dict(data_pomiaru="2022-09-01", wykonawca="GEO-B", dzialka="101/1")
_fA = domain.ingest("OP-A/2020", _hdrA, [
    dict(id="P1", x=5_500_000.00, y=6_550_000.00, **_attrs),
    dict(id="P2", x=5_500_010.00, y=6_550_010.00, **_attrs)])
check("domain: first sight all nowy",
      all(v["flaga"] == "nowy" for v in _fA.values()))
_fB = domain.ingest("OP-B/2022", _hdrB, [
    dict(id="P1", x=5_500_000.00, y=6_550_000.00, **_attrs),
    dict(id="P2", x=5_500_010.05, y=6_550_010.00, **_attrs),
    dict(id="P3", x=5_500_020.00, y=6_550_020.00, **_attrs)])
check("domain: unchanged point bez_zmian", _fB["P1"]["flaga"] == "bez_zmian")
check("domain: 5 cm move zmieniony with delta",
      _fB["P2"]["flaga"] == "zmieniony" and _fB["P2"]["delta_m"] == 0.05)
check("domain: fresh id nowy", _fB["P3"]["flaga"] == "nowy")
_rd = domain.rodowod("P2")
check("domain: supersession non-destructive",
      len(_rd) == 2 and _rd[0]["t_invalid"] and not _rd[1]["t_invalid"])
check("domain: re-ingest idempotent",
      domain.ingest("OP-B/2022", _hdrB, [
          dict(id="P1", x=5_500_000.00, y=6_550_000.00, **_attrs),
          dict(id="P2", x=5_500_010.05, y=6_550_010.00, **_attrs),
          dict(id="P3", x=5_500_020.00, y=6_550_020.00, **_attrs)]) == _fB
      and len(domain.rodowod("P2")) == 2)
domain.ingest("OP-0/2015", dict(data_pomiaru="2015-03-01", wykonawca="GEO-0"),
              [dict(id="P2", x=5_500_009.90, y=6_550_010.00, **_attrs)])
_rd2 = domain.rodowod("P2")
check("domain: backdated archive rechains history",
      [r["operat"] for r in _rd2] == ["OP-0/2015", "OP-A/2020", "OP-B/2022"]
      and _rd2[0]["flaga"] == "nowy" and _rd2[1]["flaga"] == "zmieniony"
      and _rd2[1]["delta_m"] == 0.1)
_attr_change = dict(_attrs, stb="3")
domain.ingest("OP-C/2024", dict(data_pomiaru="2024-01-05"),
              [dict(id="P1", x=5_500_000.00, y=6_550_000.00, **_attr_change)])
check("domain: attribute flip alone is zmieniony",
      domain.rodowod("P1")[-1]["flaga"] == "zmieniony")
_rj = json.load(open(os.path.join(stages.JOBS, job_name, "checks",
                                  "rodowod.json"), encoding="utf-8"))
check("domain: pipeline pins 9/9 nowy on first sight",
      len(_rj) == 9 and all(v["flaga"] == "nowy" for v in _rj.values()))
_sum = domain.summary()
check("domain: summary counts current beliefs",
      _sum["operaty"] >= 4 and _sum["punkty"] >= 12
      and sum(_sum["flagi"].values()) == _sum["punkty"]
      and _sum["flagi"].get("nowy", 0) >= 1)

# ---- rewizja: the after-run feedback loop (owner 06.08 12:38) ----
import rewizja as _rw
import store as _store
_dg = _rw.digest(job_name)
check("rewizja: digest carries kontrole/wykaz/pewnosc built by code",
      "KONTROLE:" in _dg and "9 pkt" in _dg and "PEWNO" in _dg)
_fr, _zr = _rw.wybierz_fragmenty(job_name, "co jest na s. 2?")
check("rewizja: explicit page routes that page",
      _zr == [2] and "=== STRONA 2 ===" in _fr)
_fr2, _zr2 = _rw.wybierz_fragmenty(job_name, "dlaczego punkt 152.4-03 tak?")
check("rewizja: mentioned point id brings wykaz row + its page",
      "152.4-03" in _fr2 and "PUNKTY Z WYKAZU" in _fr2 and 3 in _zr2)
_rwj = "test-rw-ch"
_rwd = os.path.join(stages.JOBS, _rwj)
os.makedirs(os.path.join(_rwd, "transcript"), exist_ok=True)
json.dump(dict(source=dict(pages=2), chapters=dict(rozdzialy=[
    dict(nr=1, tytul="Sprawozdanie", od=1, do=2)])),
    open(os.path.join(_rwd, "docmap.json"), "w", encoding="utf-8"))
open(os.path.join(_rwd, "transcript", "page-001.md"), "w",
     encoding="utf-8").write("# Strona 1\nprotokol graniczny wstep")
open(os.path.join(_rwd, "transcript", "page-002.md"), "w",
     encoding="utf-8").write("# Strona 2\nsuwnica pomiarowa szczegol")
_fr3, _zr3 = _rw.wybierz_fragmenty(_rwj, "co mowi rozdzial 1?")
check("rewizja: chapter reference routes its page range", _zr3 == [1, 2])
_fr4, _zr4 = _rw.wybierz_fragmenty(_rwj, "gdzie jest suwnica?")
check("rewizja: keyword fold-scoring finds the page", _zr4 == [2])
_calls = []
def _fake_turn(ctx_id, text, force_compact=False, user_extra=None):
    _calls.append(text)
    _store.append(ctx_id, "user", text, **(user_extra or {}))
    _store.append(ctx_id, "assistant", "ODP-" + str(len(_calls)))
    return dict(reply="ODP-" + str(len(_calls)), utilization=0.1)
_r1 = _rw.zapytaj(job_name, "co jest na s. 2?", turn_fn=_fake_turn)
_r2 = _rw.zapytaj(job_name, "a punkt 152.4-03?", turn_fn=_fake_turn)
check("rewizja: one job-bound context reused across questions",
      _r1["ctx_id"] == _r2["ctx_id"] and _r1["odpowiedz"] == "ODP-1")
_rows = _store.transcript(_r1["ctx_id"])
check("rewizja: digest note primes the context; user rows carry pytanie+zrodla",
      _rows[0].get("kind") == "digest"
      and _rows[1].get("pytanie") == "co jest na s. 2?"
      and _rows[1].get("zrodla") == [2]
      and "FRAGMENTY OPERATU" in _rows[1]["content"])
import chat as _chat
_msgs = _chat.build_messages(_rows)
check("rewizja: window system = rewident manual + digest (KV-stable prefix)",
      _msgs[0]["role"] == "system" and "REWIDENT" in _msgs[0]["content"]
      and "DIGEST OPERATU" in _msgs[0]["content"])
_an = open(os.path.join(stages.JOBS, job_name, "analizy", "rewizja.md"),
           encoding="utf-8").read()
check("rewizja: exchanges mirrored beside the documents",
      _an.count("## ") == 2 and "ODP-2" in _an and "s. 2" in _an)
_pairs = _rw.rozmowa(job_name)
check("rewizja: rozmowa pairs Q/A for the UI card",
      len(_pairs["wpisy"]) == 2
      and _pairs["wpisy"][0]["pytanie"] == "co jest na s. 2?"
      and _pairs["wpisy"][1]["odpowiedz"] == "ODP-2")
shutil.rmtree(_rwd, ignore_errors=True)
shutil.rmtree(os.path.join(_store.CONTEXTS, _r1["ctx_id"]),
              ignore_errors=True)
shutil.rmtree(os.path.join(stages.JOBS, job_name, "analizy"),
              ignore_errors=True)

# ---- nocna zmiana 2: timing, active(), civilized stop ----
def _wolny_crew(pdf):
    time.sleep(0.25)
    return dict(job="t-" + os.path.basename(pdf)[:6], PASS=True,
                werdykt="OK", flagi=0)
_sb = _ns.submit([PDF, PDF, PDF, PDF], runner=_wolny_crew)
time.sleep(0.1)
check("nocna2: active() reports the live batch", _ns.active() == _sb)
_ns.stop(_sb)
for _ in range(80):
    if _ns.status(_sb)["done"]:
        break
    time.sleep(0.1)
_sst = _ns.status(_sb)
_stat = [i["status"] for i in _sst["items"]]
check("nocna2: stop finishes current item, skips the rest",
      _sst["done"] and _sst.get("przerwana") is True
      and _stat.count("done") >= 1 and _stat.count("pominiety") >= 1
      and _stat.count("done") + _stat.count("pominiety") == 4)
check("nocna2: per-item seconds recorded for done items",
      all(i.get("sekundy", 0) >= 0.2 for i in _sst["items"]
          if i["status"] == "done"))
check("nocna2: active() clears after the batch", _ns.active() is None)
_rap2 = open(os.path.join(ROOT, "jobs", f"kolejka-{_sb}",
                          "raport_poranny.md"), encoding="utf-8").read()
check("nocna2: raport names the cut + carries times",
      "zatrzymana przez operatora" in _rap2 and "POMINIĘTY" in _rap2
      and "Czas [s]" in _rap2
      and _rap2.count(".pdf") >= 2)   # skipped-paths list for resume
shutil.rmtree(os.path.join(ROOT, "jobs", f"kolejka-{_sb}"),
              ignore_errors=True)

# ---- refleksja: proposals from the intervention trace + hot apply ----
import refleksja as _rf
_rfj = "test-refleksja"
_rfd = os.path.join(stages.JOBS, _rfj)
os.makedirs(os.path.join(_rfd, "checks"), exist_ok=True)
os.makedirs(os.path.join(_rfd, "analizy"), exist_ok=True)
open(os.path.join(_rfd, "overrides.jsonl"), "w", encoding="utf-8").write(
    json.dumps(dict(stage="naglowek", dopisek="daty tylko ze sprawozdania"))
    + "\n")
open(os.path.join(_rfd, "analizy", "rewizja.md"), "w", encoding="utf-8").write(
    "## pytanie\nczemu data z okladki?\n")
json.dump(dict(PASS=False, flagi=dict(wsp_jeden_silnik=[dict(id="P1")])),
          open(os.path.join(_rfd, "checks", "invariants.json"), "w",
               encoding="utf-8"))
_mat = _rf.material(_rfj)
check("refleksja: material carries dopiski + rewizja + flagi",
      "daty tylko ze sprawozdania" in _mat and "czemu data z okladki" in _mat
      and "wsp_jeden_silnik" in _mat)
_prop = _rf.zaproponuj(_rfj, runner=lambda payload: dict(propozycje=[
    dict(agent="ekstraktor_naglowka",
         zasady=[dict(typ="pozytywna", tresc="Daty bierz ze sprawozdania.")],
         uzasadnienie="dopisek operatora"),
    dict(agent="nieznany_agent", zasady=[dict(typ="negatywna", tresc="x")],
         uzasadnienie="szum")]))
check("refleksja: unknown agents filtered out",
      len(_prop["propozycje"]) == 1
      and _prop["propozycje"][0]["agent"] == "ekstraktor_naglowka")
_przed = llm.prompt("ekstraktor_naglowka")
_info = _rf.zastosuj("ekstraktor_naglowka",
                     [dict(typ="pozytywna", tresc="Daty bierz ze sprawozdania.")])
_po = llm.prompt("ekstraktor_naglowka")
check("refleksja: apply writes user overlay + manual hot-reloads (no restart)",
      _info["source"] == "user" and "Daty bierz ze sprawozdania." in _po
      and "# DOPISKI OPERATORA" in _po and _po.startswith(_przed[:200]))
os.remove(os.path.join(ROOT, "prompts", "user", "ekstraktor_naglowka.md"))
check("refleksja: overlay removal restores golden",
      llm.prompt("ekstraktor_naglowka") == _przed)
shutil.rmtree(_rfd, ignore_errors=True)

# ---- rozmowa: conversation-first layer (owner 16:37) ----
import rozmowa as _rz2
_scen = iter([
    dict(ma_wykaz_granic=True, strony_wykazu_kandydaci=[3],
         strony_sprawozdania_kandydaci=[2]),
    dict(id_zgloszenia="GK.1/26", dzialki=["152/4"], uklad="PL-2000/6",
         wykonawca="GEO-X"),
    dict(liczba_punktow=9, strona_wykazu=3, punkty=[]),
    dict(werdykt="OK", zaufanie="wysokie", problemy=[]),
])
def _rz_runner(job, st, dopisek=None):
    return next(_scen)
_ro = _rz2.start(PDF, "t-rozmowa")
_ctx2 = _ro["rozmowa"]
check("rozmowa: start binds job + opens with file drop and welcome",
      _ro["next_stage"] == "klasyfikator"
      and [r.get("kind") for r in _store.transcript(_ctx2)]
      == ["digest", "plik", "powitanie"])
for _oczek in ("klasyfikator", "naglowek", "wykaz"):
    _r = _rz2.krok(_ctx2, "run", runner=_rz_runner)
    check(f"rozmowa: stage {_oczek} narrated into the SAME context",
          _r["awaiting"] and _r["stage"] == _oczek
          and _store.transcript(_ctx2)[-1].get("kind") == "etap")
    _rz2.krok(_ctx2, "accept")
os.makedirs(os.path.join(stages.JOBS, _ro["job"], "out"), exist_ok=True)
open(os.path.join(stages.JOBS, _ro["job"], "out", "przeglad.csv"),
     "w").write("x")
json.dump(dict(PASS=True), open(os.path.join(
    stages.JOBS, _ro["job"], "checks", "invariants.json"), "w"))
_r = _rz2.krok(_ctx2, "run", runner=_rz_runner)      # werdykt
_r2 = _rz2.krok(_ctx2, "accept")                     # -> done/final
_rows2 = _store.transcript(_ctx2)
_fin = [r for r in _rows2 if r.get("kind") == "final"]
check("rozmowa: completion message carries file attachments + PASS",
      _r2.get("done") and len(_fin) == 1
      and _fin[0]["zalaczniki"][0]["plik"] == "przeglad.csv"
      and _fin[0]["pass_"] is True
      and _rows2[-1].get("kind") == "digest")
_rz2.krok(_ctx2, "accept")
check("rozmowa: marker lifecycle + validation",
      _rz2.marker(_ctx2, "zatwierdzona")["marker"] == "zatwierdzona"
      and _store.meta(_ctx2)["marker"] == "zatwierdzona")
try:
    _rz2.marker(_ctx2, "xxx"); _zle = False
except ValueError:
    _zle = True
check("rozmowa: unknown marker rejected", _zle)
# Legacy-job fixture: jobs/ jest nietrackowane, wiec swiezy checkout (CI)
# nie ma zadnego legacy joba do wylistowania. Tworzymy minimalny ksztalt
# tylko gdy go brak i sprzatamy dokladnie to, co stworzylismy.
_legacy_job = "job-20260806-133900"
_legacy_dir = os.path.join(stages.JOBS, _legacy_job)
_kolejka_ci = os.path.join(stages.JOBS, "kolejka-ci")
_made_legacy = not os.path.isdir(_legacy_dir)
if _made_legacy:
    os.makedirs(os.path.join(_legacy_dir, "extracted"), exist_ok=True)
    json.dump(dict(dzialka="140/2"), open(os.path.join(
        _legacy_dir, "extracted", "header.json"), "w", encoding="utf-8"))
    os.makedirs(_kolejka_ci, exist_ok=True)
    json.dump(dict(items=[dict(job=_legacy_job)]), open(os.path.join(
        _kolejka_ci, "kolejka.json"), "w", encoding="utf-8"))
_ls = _rz2.lista()
_moja = [r for r in _ls if r["id"] == _ctx2][0]
check("rozmowa: rail entry is document-named with machine+human status",
      _moja["nazwa"].startswith("operat_sample.pdf")
      and _moja["marker"] == "zatwierdzona" and _moja["maszyna"] is True)
_legacy = [r for r in _ls if r["job"] == "job-20260806-133900"]
check("rozmowa: legacy job named by dzialka, batch tagged with kolejka",
      _legacy and _legacy[0]["nazwa"].startswith("dz. ")
      and _legacy[0].get("kolejka"))
if _made_legacy:
    shutil.rmtree(_legacy_dir, ignore_errors=True)
    shutil.rmtree(_kolejka_ci, ignore_errors=True)
def _fake_turn2(cid, text, force_compact=False, user_extra=None):
    _store.append(cid, "user", text, **(user_extra or {}))
    _store.append(cid, "assistant", "ODP-R")
    return dict(reply="ODP-R", utilization=0.1)
_pq = _rz2.pytanie(_ctx2, "co jest na s. 2?", turn_fn=_fake_turn2)
check("rozmowa: question rides the SAME conversation context",
      _pq["ctx_id"] == _ctx2
      and _store.transcript(_ctx2)[-1]["content"] == "ODP-R")
shutil.rmtree(os.path.join(stages.JOBS, _ro["job"]), ignore_errors=True)
shutil.rmtree(os.path.join(_store.CONTEXTS, _ctx2), ignore_errors=True)

# ---- dedup punktow: powtorzenie != blad, sprzecznosc = flaga (masywny TP) ----
_dp = [dict(id="A", x=100.00, y=200.00, page=3),
       dict(id="A", x=100.005, y=200.00, page=5),   # zgodna kopia -> merge
       dict(id="B", x=300.00, y=400.00, page=3),
       dict(id="B", x=900.00, y=400.00, page=5),    # sprzecznosc -> zostaje
       dict(id="C", x=1.00, y=2.00, page=4)]
_dd, _du = stages.dedup_punkty(_dp)
check("dedup: agreeing copy merged with provenance note",
      len(_dd) == 4 and len([p for p in _dd if p["id"] == "A"]) == 1
      and any("scalono" in u for u in _du))
check("dedup: conflicting copies both kept (truth law)",
      len([p for p in _dd if p["id"] == "B"]) == 2)
_chk_s = dict(flagi={"sprzeczne_id": [dict(id="B", strony=[3, 5],
                                           delta_m=600.0)]})
_pw = stages.pewnosc_punktow(_dd, _chk_s)
check("dedup: contradiction paints the point KRYTYCZNE with the reason",
      _pw["B"]["tier"] == "KRYTYCZNE"
      and any("sprzecznymi" in p for p in _pw["B"]["powody"])
      and _pw["C"]["tier"] == "PEWNE")

# ---- live gates: follow_log streams and terminates on pieczec ----
import server as _srv
_evt = list(_srv.follow_log(job_name, timeout_s=5))
check("live: follow_log replays the whole run and stops at pieczec",
      len(_evt) >= 6 and '"stage": "pieczec"' in _evt[-1]
      and '"stage": "ingest"' in _evt[0])

# ---- launcher model-pack integrity (size gate + one-time deep hash) ----
import hashlib
import importlib.util as _ilu
_spec = _ilu.spec_from_file_location(
    "start_geodoc", os.path.join(ROOT, "start_geodoc.pyw"))
_lch = _ilu.module_from_spec(_spec)
_spec.loader.exec_module(_lch)
_tmp = os.path.join(ROOT, "jobs", "pack-test")
shutil.rmtree(_tmp, ignore_errors=True)
os.makedirs(os.path.join(_tmp, "models"))
_payload = b"GGUFfake" * 100
with open(os.path.join(_tmp, "models", "m.gguf"), "wb") as f:
    f.write(_payload)
_man = dict(wersja=1, pliki=[dict(name="m.gguf", bytes=len(_payload),
                                  sha256=hashlib.sha256(_payload).hexdigest())])
with open(os.path.join(_tmp, "models", "manifest.json"), "w") as f:
    json.dump(_man, f)
_lch.ROOT = _tmp
check("pack: deep verify passes + writes marker", _lch.verify_models() is True
      and os.path.isfile(os.path.join(_tmp, "models", ".verified")))
with open(os.path.join(_tmp, "models", "m.gguf"), "ab") as f:
    f.write(b"X")
check("pack: size drift fails instantly", _lch.verify_models() is False
      and os.path.isfile(os.path.join(_tmp, "geodoc-error.log")))
shutil.rmtree(_tmp, ignore_errors=True)

# ---- self-description: the README's measured counts match their sources ----
# The README states two counts; their homes are this harness and the arsenal
# manifest. These checks keep the sentence and the source from drifting.
import re as _re  # noqa: E402
_readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
_arsenal = open(os.path.join(ROOT, "testdata", "ARSENAL.md"), encoding="utf-8").read()
_rows = [r for r in _arsenal.splitlines() if _re.match(r"^\| \**operat_\w+", r)]
_named = [_re.match(r"^\| \**(operat_\w+)", r).group(1) for r in _rows]
check("readme: every arsenal fixture exists in testdata/",
      all(os.path.isfile(os.path.join(ROOT, "testdata", n + ".pdf")) for n in _named),
      str([n for n in _named
           if not os.path.isfile(os.path.join(ROOT, "testdata", n + ".pdf"))]))
_m_ars = _re.search(r"\| (\d+) synthetic documents built to break it", _readme)
check("readme: adversarial-document count = rows in testdata/ARSENAL.md",
      bool(_m_ars) and int(_m_ars.group(1)) == len(_rows),
      f"README {_m_ars.group(1) if _m_ars else None} vs ARSENAL {len(_rows)}")
# last check on purpose: it counts every check including itself
_m_chk = _re.search(r"\| Check suite \| (\d+) checks", _readme)
check("readme: check count = checks this harness runs",
      bool(_m_chk) and int(_m_chk.group(1)) == len(CHECKS) + 1,
      f"README {_m_chk.group(1) if _m_chk else None} vs harness {len(CHECKS) + 1}")

print(f"\n{'ALL GREEN' if not FAILS else 'FAILURES: ' + ', '.join(FAILS)} "
      f"({len(FAILS)} fail)")
sys.exit(1 if FAILS else 0)
